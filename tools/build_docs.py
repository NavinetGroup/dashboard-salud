# -*- coding: utf-8 -*-
"""Build the Word (.docx) versions of the manuals from their Markdown sources.

    python tools/build_docs.py

Reads docs/MANUAL_USUARIO.md and docs/MANUAL_TECNICO.md and writes the .docx
next to each. The Markdown files are the source of truth; edit them and rerun.
Supports the Markdown this project uses: headings, paragraphs, bullet and
numbered lists (nested by indentation), pipe tables, fenced code blocks, and
inline **bold**, `code` and [links](url). Needs python-docx (in
requirements-pipeline.txt).
"""

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / 'docs'
MANUALS = [
    ('MANUAL_USUARIO.md', 'Manual de usuario', 'Tablero Informe Regional de Salud'),
    ('MANUAL_TECNICO.md', 'Manual técnico', 'Fuentes, procesamiento, tablero y distribución'),
]
TEAL = RGBColor(0x0B, 0x5D, 0x5E)
_INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def _shade(cell, hex_fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_fill)
    pr.append(shd)


def _runs(par, text: str, mono_size: int = 9) -> None:
    """Add text to a paragraph, rendering the inline Markdown marks."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = par.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(mono_size)
        elif part.startswith('['):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part).groups()
            par.add_run(label).bold = True
            if url.startswith('http'):
                par.add_run(f' ({url})').italic = True
        else:
            par.add_run(part)


def _toc(doc) -> None:
    par = doc.add_paragraph()
    run = par.add_run()
    for tag, text in (('begin', None), ('instr', 'TOC \\o "1-3" \\h \\z \\u'), ('separate', None),
                      ('text', 'Tabla de contenido: pulse F9 (o "Actualizar campo") para generarla.'),
                      ('end', None)):
        if tag == 'instr':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        elif tag == 'text':
            el = OxmlElement('w:t')
            el.text = text
        else:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        run._r.append(el)
    # ask Word to refresh fields (the TOC) when the file is opened
    settings = doc.settings.element
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    settings.append(upd)


def _table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ''
            par = cell.paragraphs[0]
            _runs(par, text, mono_size=8)
            for r in par.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
            if i == 0:
                _shade(cell, 'E0F2F1')
    doc.add_paragraph()


def _code(doc, lines: list[str]) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(12)
    for k, line in enumerate(lines):
        r = par.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        if k < len(lines) - 1:
            r.add_break(WD_BREAK.LINE)
    _shade_par(par)


def _shade_par(par) -> None:
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F4F6F7')
    ppr.append(shd)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip('|').split('|')]


def convert(md_path: Path, title: str, subtitle: str) -> Path:
    doc = Document()
    styles = doc.styles
    styles['Normal'].font.name = 'Calibri'
    styles['Normal'].font.size = Pt(10.5)
    for lvl in (1, 2, 3):
        styles[f'Heading {lvl}'].font.color.rgb = TEAL

    # Cover
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run('Informe Regional de Salud')
    r.font.size, r.font.color.rgb, r.bold = Pt(26), TEAL, True
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.size, r.bold = Pt(20), True
    doc.add_paragraph(subtitle).runs[0].font.size = Pt(13)
    doc.add_paragraph(f'Supersalud — versión del {date.today():%d/%m/%Y}').runs[0].italic = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _toc(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    lines = md_path.read_text(encoding='utf-8').splitlines()
    i, para, first_h1 = 0, [], True

    def flush():
        if para:
            _runs(doc.add_paragraph(), ' '.join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i])
                i += 1
            _code(doc, block)
        elif line.startswith('#'):
            flush()
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            if level == 1 and first_h1:      # the document title is on the cover already
                first_h1 = False
            else:
                doc.add_heading(text, level=min(level, 4) if level > 1 else 1)
        elif line.strip().startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not re.match(r'^\s*\|[\s\-:|]+\|\s*$', lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            _table(doc, [r + [''] * (width - len(r)) for r in rows])
            continue
        elif re.match(r'^\s*([-*]|\d+\.)\s', line):
            flush()
            indent = (len(line) - len(line.lstrip(' '))) // 2
            m = re.match(r'^\s*(?:[-*]|(\d+)\.)\s+(.*)$', line)
            style = 'List Number' if m.group(1) else 'List Bullet'
            if indent:
                style += f' {min(indent + 1, 3)}'
            # a wrapped list item continues on indented lines
            text = m.group(2)
            while i + 1 < len(lines) and lines[i + 1].startswith('  ') and \
                    not re.match(r'^\s*([-*]|\d+\.)\s', lines[i + 1]) and lines[i + 1].strip():
                i += 1
                text += ' ' + lines[i].strip()
            _runs(doc.add_paragraph(style=style), text)
        elif line.strip() in ('', '---'):
            flush()
        elif line.startswith('>'):
            flush()
            p = doc.add_paragraph(style='Intense Quote')
            _runs(p, line.lstrip('> ').strip())
        else:
            para.append(line)
        i += 1
    flush()
    out = md_path.with_suffix('.docx')
    doc.save(out)
    return out


def main() -> int:
    for name, title, subtitle in MANUALS:
        src = DOCS / name
        if not src.exists():
            print(f'falta {src}')
            return 2
        out = convert(src, title, subtitle)
        print(f'{out.relative_to(ROOT)} ({out.stat().st_size / 1e3:.0f} KB)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
